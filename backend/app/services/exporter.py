"""Tabular exports to Excel (RTL), CSV (UTF-8 BOM for Excel) and PDF (Persian-shaped text)."""
import csv
import io
from collections.abc import Sequence
from datetime import date, datetime
from urllib.parse import quote

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config.settings import settings
from app.core.exceptions import ValidationFailed
from app.utils.jalali import format_jalali

Column = tuple[str, str]  # (key, Persian header)

_MIME = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "pdf": "application/pdf",
}


def _cell(value):
    if isinstance(value, datetime):
        return format_jalali(value)
    if isinstance(value, date):
        return format_jalali(value, with_time=False)
    if isinstance(value, bool):
        return "بله" if value else "خیر"
    if isinstance(value, (list, tuple)):
        return "، ".join(str(v) for v in value)
    return "" if value is None else value


def to_xlsx(title: str, columns: Sequence[Column], rows: Sequence[dict], sheets: dict | None = None) -> bytes:
    wb = Workbook()
    all_sheets = {title: (columns, rows), **(sheets or {})}
    first = True
    for sheet_title, (cols, data) in all_sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = sheet_title[:31]
        ws.sheet_view.rightToLeft = True
        header_fill = PatternFill("solid", fgColor="EEF2FF")
        thin = Side(style="thin", color="E5E7EB")
        for idx, (_, header) in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=idx, value=header)
            cell.font = Font(bold=True, name="Vazirmatn")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin)
        widths = [len(h) + 4 for _, h in cols]
        for r_idx, row in enumerate(data, start=2):
            for c_idx, (key, _) in enumerate(cols, start=1):
                value = _cell(row.get(key))
                ws.cell(row=r_idx, column=c_idx, value=value).alignment = Alignment(horizontal="right")
                widths[c_idx - 1] = min(60, max(widths[c_idx - 1], len(str(value)) + 2))
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_csv(columns: Sequence[Column], rows: Sequence[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([h for _, h in columns])
    for row in rows:
        writer.writerow([_cell(row.get(k)) for k, _ in columns])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def _shape(text) -> str:
    import arabic_reshaper
    from bidi.algorithm import get_display

    return get_display(arabic_reshaper.reshape(str(text)))


def to_pdf(title: str, columns: Sequence[Column], rows: Sequence[dict], subtitle: str | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = "Helvetica", "Helvetica-Bold"
    if settings.PDF_FONT_PATH.exists():
        if "Vazirmatn" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Vazirmatn", str(settings.PDF_FONT_PATH)))
            bold_path = settings.PDF_FONT_BOLD_PATH if settings.PDF_FONT_BOLD_PATH.exists() else settings.PDF_FONT_PATH
            pdfmetrics.registerFont(TTFont("Vazirmatn-Bold", str(bold_path)))
        font, bold = "Vazirmatn", "Vazirmatn-Bold"

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1.2 * cm, leftMargin=1.2 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm, title=title)
    title_style = ParagraphStyle("t", fontName=bold, fontSize=15, alignment=2, leading=22)
    sub_style = ParagraphStyle("s", fontName=font, fontSize=9, alignment=2, textColor=colors.HexColor("#6B7280"))
    cell_style = ParagraphStyle("c", fontName=font, fontSize=8, alignment=2, leading=12)
    head_style = ParagraphStyle("h", fontName=bold, fontSize=8.5, alignment=1, leading=12)

    story = [Paragraph(_shape(title), title_style)]
    story.append(Paragraph(_shape(subtitle or f"تاریخ تهیه گزارش: {format_jalali(datetime.utcnow())}"), sub_style))
    story.append(Spacer(1, 10))
    # RTL table: reverse column order so the first column is on the right
    cols = list(reversed(columns))
    data = [[Paragraph(_shape(h), head_style) for _, h in cols]]
    for row in rows[:5000]:
        data.append([Paragraph(_shape(_cell(row.get(k))), cell_style) for k, _ in cols])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2FF")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E7EB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
    ]))
    story.append(table)
    doc.build(story)
    return buffer.getvalue()


def export_response(fmt: str, filename: str, title: str, columns: Sequence[Column], rows: Sequence[dict],
                    subtitle: str | None = None, sheets: dict | None = None) -> StreamingResponse:
    if fmt == "xlsx":
        content = to_xlsx(title, columns, rows, sheets)
    elif fmt == "csv":
        content = to_csv(columns, rows)
    elif fmt == "pdf":
        content = to_pdf(title, columns, rows, subtitle)
    else:
        raise ValidationFailed("فرمت خروجی پشتیبانی نمی‌شود.")
    stamp = datetime.utcnow().strftime("%Y%m%d-%H%M")
    name = f"{filename}-{stamp}.{fmt}"
    return StreamingResponse(
        io.BytesIO(content), media_type=_MIME[fmt],
        headers={"Content-Disposition": f"attachment; filename=\"{name}\"; filename*=UTF-8''{quote(name)}"},
    )


TICKET_COLUMNS: list[Column] = [
    ("code", "شماره تیکت"), ("subject", "موضوع"), ("customer", "مشتری"), ("customer_contact", "اطلاعات تماس"),
    ("status", "وضعیت"), ("priority", "اولویت"), ("department", "دپارتمان"), ("category", "دسته‌بندی"),
    ("agent", "کارشناس"), ("tags", "برچسب‌ها"), ("sla_status", "وضعیت SLA"), ("created_at", "تاریخ ایجاد"),
    ("last_response_at", "آخرین پاسخ"), ("resolved_at", "تاریخ حل"), ("closed_at", "تاریخ بستن"),
]

SLA_LABELS = {"none": "بدون SLA", "healthy": "در مهلت", "warning": "هشدار", "breached": "نقض شده",
              "paused": "متوقف", "met": "رعایت شده"}


def ticket_rows(tickets) -> list[dict]:
    return [{
        "code": t.code, "subject": t.subject, "customer": t.customer.full_name if t.customer else "",
        "customer_contact": (t.customer.mobile or t.customer.email) if t.customer else "",
        "status": t.status.name, "priority": t.priority.name,
        "department": t.department.name if t.department else "", "category": t.category.name if t.category else "",
        "agent": t.assigned_agent.full_name if t.assigned_agent else "", "tags": [tag.name for tag in t.tags],
        "sla_status": SLA_LABELS.get(t.sla_status, t.sla_status), "created_at": t.created_at,
        "last_response_at": t.last_response_at, "resolved_at": t.resolved_at, "closed_at": t.closed_at,
    } for t in tickets]


USER_COLUMNS: list[Column] = [
    ("full_name", "نام و نام خانوادگی"), ("email", "ایمیل"), ("mobile", "موبایل"), ("organization", "سازمان/شرکت"),
    ("job_title", "سمت"), ("roles", "نقش‌ها"), ("departments", "دپارتمان‌ها"), ("is_active", "فعال"),
    ("last_login_at", "آخرین ورود"), ("created_at", "تاریخ ثبت"),
]


def user_rows(users) -> list[dict]:
    return [{
        "full_name": u.full_name, "email": u.email, "mobile": u.mobile, "organization": u.organization,
        "job_title": u.job_title, "roles": [r.display_name for r in u.roles],
        "departments": [d.name for d in u.departments], "is_active": u.is_active,
        "last_login_at": u.last_login_at, "created_at": u.created_at,
    } for u in users]
