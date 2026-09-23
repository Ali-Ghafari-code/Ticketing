"""Excel import for customers, users and categories with template, validation, preview and duplicate detection."""
import io
import secrets
from dataclasses import dataclass, field

from email_validator import EmailNotValidError, validate_email
from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.exceptions import ValidationFailed
from app.core.security import hash_password
from app.database.base import utcnow
from app.models import Company, Department, Role, TicketCategory, User, UserType
from app.services.company import ensure_agent_capacity, ensure_customer_capacity
from app.services.rbac import assignable_roles_query, get_system_role
from app.utils.persian import normalize_mobile, normalize_text

MAX_ROWS = 5000

TEMPLATES = {
    "customers": {
        "title": "مشتریان",
        "columns": [("full_name", "نام و نام خانوادگی *"), ("mobile", "موبایل"), ("email", "ایمیل"),
                    ("organization", "شرکت/سازمان"), ("notes", "توضیحات")],
        "sample": ["علی رضایی", "09121234567", "ali@example.com", "شرکت نمونه", ""],
    },
    "users": {
        "title": "کاربران",
        "columns": [("full_name", "نام و نام خانوادگی *"), ("mobile", "موبایل"), ("email", "ایمیل"),
                    ("role", "نقش * (company_admin / support_manager / support_agent)"),
                    ("department", "دپارتمان"), ("job_title", "سمت")],
        "sample": ["مریم احمدی", "09351234567", "maryam@example.com", "support_agent", "پشتیبانی فنی", "کارشناس"],
    },
    "categories": {
        "title": "دسته‌بندی‌ها",
        "columns": [("name", "نام دسته‌بندی *"), ("parent", "دسته‌بندی والد"), ("department", "دپارتمان"),
                    ("description", "توضیحات")],
        "sample": ["مشکل ورود", "فنی", "پشتیبانی فنی", "مشکلات مربوط به ورود به حساب"],
    },
}


def template_xlsx(kind: str) -> bytes:
    spec = TEMPLATES.get(kind)
    if spec is None:
        raise ValidationFailed("نوع قالب نامعتبر است.")
    wb = Workbook()
    ws = wb.active
    ws.title = spec["title"]
    ws.sheet_view.rightToLeft = True
    for idx, (_, header) in enumerate(spec["columns"], start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF2FF")
        ws.column_dimensions[cell.column_letter].width = max(18, len(header) + 4)
    for idx, value in enumerate(spec["sample"], start=1):
        ws.cell(row=2, column=idx, value=value)
    if kind == "users":
        dv = DataValidation(type="list", formula1='"company_admin,support_manager,support_agent"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add("D2:D5000")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@dataclass
class ImportRow:
    row: int
    data: dict
    errors: list[str] = field(default_factory=list)
    duplicate: bool = False

    @property
    def valid(self) -> bool:
        return not self.errors


def _read(upload: UploadFile, kind: str) -> list[dict]:
    if not (upload.filename or "").lower().endswith(".xlsx"):
        raise ValidationFailed("فقط فایل اکسل با پسوند xlsx پذیرفته می‌شود.")
    content = upload.file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise ValidationFailed("حجم فایل اکسل نباید بیش از ۱۰ مگابایت باشد.")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types
        raise ValidationFailed("فایل اکسل قابل خواندن نیست.") from exc
    ws = wb.worksheets[0]
    keys = [k for k, _ in TEMPLATES[kind]["columns"]]
    rows = []
    for idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if idx - 1 > MAX_ROWS:
            raise ValidationFailed(f"حداکثر {MAX_ROWS} ردیف در هر بار ورود مجاز است.")
        values = list(values or [])
        if not any(v not in (None, "") for v in values):
            continue
        rows.append({"_row": idx, **{k: (str(values[i]).strip() if i < len(values) and values[i] is not None else "")
                                     for i, k in enumerate(keys)}})
    return rows


def _contact_checks(db: Session, item: ImportRow, seen_emails: set, seen_mobiles: set) -> None:
    data = item.data
    if data.get("mobile"):
        try:
            data["mobile"] = normalize_mobile(data["mobile"])
        except ValueError as exc:
            item.errors.append(str(exc))
    if data.get("email"):
        try:
            data["email"] = validate_email(data["email"], check_deliverability=False).normalized.lower()
        except EmailNotValidError:
            item.errors.append("ایمیل معتبر نیست.")
    if not data.get("email") and not data.get("mobile"):
        item.errors.append("ایمیل یا موبایل الزامی است.")
    for key, seen in (("email", seen_emails), ("mobile", seen_mobiles)):
        value = data.get(key)
        if not value:
            continue
        if value in seen:
            item.duplicate = True
            item.errors.append(f"{'ایمیل' if key == 'email' else 'موبایل'} در فایل تکراری است.")
        seen.add(value)
        if db.scalar(select(User.id).where(getattr(User, key) == value)):
            item.duplicate = True
            item.errors.append(f"{'ایمیل' if key == 'email' else 'موبایل'} قبلاً در سامانه ثبت شده است.")


def validate_rows(db: Session, company: Company, kind: str, rows: list[dict]) -> list[ImportRow]:
    result: list[ImportRow] = []
    seen_emails: set[str] = set()
    seen_mobiles: set[str] = set()
    departments = {normalize_text(d.name): d for d in db.scalars(
        select(Department).where(Department.company_id == company.id, Department.deleted_at.is_(None)))}
    categories = {normalize_text(c.name): c for c in db.scalars(
        select(TicketCategory).where(TicketCategory.company_id == company.id, TicketCategory.deleted_at.is_(None)))}
    roles = {r.name: r for r in db.scalars(assignable_roles_query(company.id))}
    seen_categories: set[str] = set()

    for raw in rows:
        row_no = raw.pop("_row", 0)
        item = ImportRow(row=row_no, data=dict(raw))
        data = item.data
        if kind in ("customers", "users"):
            if len(data.get("full_name", "")) < 2:
                item.errors.append("نام و نام خانوادگی الزامی است.")
            _contact_checks(db, item, seen_emails, seen_mobiles)
        if kind == "users":
            role = roles.get(data.get("role", ""))
            if role is None or role.audience != "staff":
                item.errors.append("نقش وارد شده معتبر نیست.")
            if data.get("department") and normalize_text(data["department"]) not in departments:
                item.errors.append("دپارتمان یافت نشد.")
        if kind == "categories":
            name = normalize_text(data.get("name", ""))
            if not name or len(name) < 2:
                item.errors.append("نام دسته‌بندی الزامی است.")
            elif name in categories or name in seen_categories:
                item.duplicate = True
                item.errors.append("این دسته‌بندی قبلاً وجود دارد.")
            seen_categories.add(name or "")
            parent = normalize_text(data.get("parent", ""))
            if parent and parent not in categories and parent not in seen_categories:
                item.errors.append("دسته‌بندی والد یافت نشد (والد باید قبل از فرزند بیاید).")
            if data.get("department") and normalize_text(data["department"]) not in departments:
                item.errors.append("دپارتمان یافت نشد.")
        result.append(item)
    return result


def preview(db: Session, company: Company, kind: str, upload: UploadFile) -> dict:
    if kind not in TEMPLATES:
        raise ValidationFailed("نوع ورود اطلاعات نامعتبر است.")
    rows = validate_rows(db, company, kind, _read(upload, kind))
    return {
        "kind": kind,
        "total": len(rows),
        "valid": sum(1 for r in rows if r.valid),
        "invalid": sum(1 for r in rows if not r.valid),
        "duplicates": sum(1 for r in rows if r.duplicate),
        "columns": [{"key": k, "label": label} for k, label in TEMPLATES[kind]["columns"]],
        "rows": [{"row": r.row, "data": r.data, "errors": r.errors, "duplicate": r.duplicate, "valid": r.valid}
                 for r in rows],
    }


def commit(db: Session, company: Company, kind: str, rows: list[dict]) -> dict:
    """Re-validates the rows from the preview (never trusts the client) and inserts the valid ones."""
    if kind not in TEMPLATES:
        raise ValidationFailed("نوع ورود اطلاعات نامعتبر است.")
    if len(rows) > MAX_ROWS:
        raise ValidationFailed(f"حداکثر {MAX_ROWS} ردیف در هر بار ورود مجاز است.")
    allowed = {k for k, _ in TEMPLATES[kind]["columns"]}
    cleaned = [{"_row": r.get("row", i + 2), **{k: str(v or "").strip() for k, v in (r.get("data") or {}).items()
                                                  if k in allowed}} for i, r in enumerate(rows)]
    validated = validate_rows(db, company, kind, cleaned)
    departments = {normalize_text(d.name): d for d in db.scalars(
        select(Department).where(Department.company_id == company.id, Department.deleted_at.is_(None)))}
    created = 0
    if kind == "customers":
        customer_role = get_system_role(db, "customer")
        for item in validated:
            if not item.valid:
                continue
            ensure_customer_capacity(db, company)
            d = item.data
            user = User(company_id=company.id, user_type=UserType.CUSTOMER, full_name=d["full_name"],
                        email=d.get("email") or None, mobile=d.get("mobile") or None,
                        organization=d.get("organization") or None, notes=d.get("notes") or None,
                        password_hash=hash_password(secrets.token_urlsafe(16)))
            user.roles = [customer_role]
            db.add(user)
            db.flush()
            created += 1
    elif kind == "users":
        roles = {r.name: r for r in db.scalars(assignable_roles_query(company.id))}
        for item in validated:
            if not item.valid:
                continue
            ensure_agent_capacity(db, company)
            d = item.data
            user = User(company_id=company.id, user_type=UserType.STAFF, full_name=d["full_name"],
                        email=d.get("email") or None, mobile=d.get("mobile") or None,
                        job_title=d.get("job_title") or None, password_hash=hash_password(secrets.token_urlsafe(16)),
                        password_changed_at=utcnow())
            user.roles = [roles[d["role"]]]
            if d.get("department"):
                user.departments = [departments[normalize_text(d["department"])]]
            db.add(user)
            db.flush()
            created += 1
    else:
        existing = {normalize_text(c.name): c for c in db.scalars(
            select(TicketCategory).where(TicketCategory.company_id == company.id, TicketCategory.deleted_at.is_(None)))}
        for item in validated:
            if not item.valid:
                continue
            d = item.data
            parent = existing.get(normalize_text(d.get("parent", ""))) if d.get("parent") else None
            dept = departments.get(normalize_text(d.get("department", ""))) if d.get("department") else None
            category = TicketCategory(company_id=company.id, name=d["name"], description=d.get("description") or None,
                                      parent_id=parent.id if parent else None, department_id=dept.id if dept else None)
            db.add(category)
            db.flush()
            existing[normalize_text(d["name"])] = category
            created += 1
    skipped = len(validated) - created
    return {"created": created, "skipped": skipped,
            "errors": [{"row": r.row, "errors": r.errors} for r in validated if not r.valid]}


__all__ = ["template_xlsx", "preview", "commit", "TEMPLATES", "Role"]
